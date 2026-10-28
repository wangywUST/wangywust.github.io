"""Exercise section binding, preservation, repeated updates, and locked files."""

from io import BytesIO
from pathlib import Path
import json
import os
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from faculty_form_word import update_faculty_form


class FacultyFormTests(unittest.TestCase):
    def fixture(self, folder):
        root = Path(folder)
        data = root / "data"
        data.mkdir()
        config = json.loads((ROOT / "data/faculty_form.json").read_text(encoding="utf-8"))
        (data / "faculty_form.json").write_text(json.dumps(config), encoding="utf-8")
        records = {
            "profile": {"name_zh": "王老师", "website": "https://example.test/"},
            "education": {"items": [{"primary_zh": "某大学，本科", "date_zh": "2013年 - 2017年"}, {"primary_zh": "某大学，博士", "date_zh": "2019年 - 2023年"}]},
            "work": {"items": [{"primary_zh": "某大学，博士后研究员", "date_zh": "2023年 - 2024年"}, {"primary_zh": "某大学，教授", "date_zh": "2025年 - 至今"}]},
            "teaching": {"items": [{"primary_zh": "某大学，CSE 188 主讲教师", "date": "2026"}, {"primary_zh": "某大学，编程课程助教", "date": "2020"}]},
            "grants": {"items": [{"primary_zh": "新增基金项目", "date": "2026"}]},
            "awards": {"items": [{"primary_zh": "国家奖学金", "date": "2015"}, {"primary_zh": "论文奖", "date": "2026"}]},
            "editorship": {"items": [{"primary": "Transactions on Machine Learning Research (TMLR)"}, {"primary": "Pattern Recognition", "ratings": ["JCR Q1"]}]},
            "area_chair": {"items": [{"primary": "Conference on Empirical Methods in Natural Language Processing (EMNLP)"}, {"primary": "The Web Conference (WWW)", "ratings": ["CCF-A"]}]},
            "tutorial": {"items": [{"primary_zh": "世界模型", "primary_en": "World Models", "date": "CVPR 2026"}]},
            "reviewing": {"items": [{"primary_zh": "会议审稿人", "secondary": "ACL; WWW"}, {"primary_zh": "期刊审稿人", "secondary": "Journal A; Journal B"}]},
        }
        for key, value in records.items():
            (data / f"{key}.json").write_text(json.dumps(value), encoding="utf-8")
        document = Document()
        table = document.add_table(rows=0, cols=1)
        table.style = "Table Grid"
        def section(title, lines):
            cell = table.add_row().cells[0]
            run = cell.paragraphs[0].add_run(title)
            run.bold, run.font.size, run.font.color.rgb = True, Pt(12), RGBColor(255, 255, 255)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "365F91")
            cell._tc.get_or_add_tcPr().append(shading)
            if lines is not None:
                cell = table.add_row().cells[0]
                for index, text in enumerate(lines):
                    p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
                    r = p.add_run(text)
                    r.bold, r.font.size = True, Pt(12)
        section("学习经历：", ["手写的详细学习经历和项目说明，必须保留。"])
        section("工作经历：", [""])
        # No teaching header: the renderer must insert it before grants.
        section("主持科研项目：", ["旧项目"])
        section("曾获奖项：", ["本科期间：", "2015年，旧奖项", "任教期间：", "2026年，旧论文奖"])
        section("专业活动：", ["期刊编委：", "Transactions on Machine Learning Research (TMLR)，担任 Action Editor。", "会议领域主席：", "Conference on Empirical Methods in Natural Language Processing (EMNLP)，担任 Senior Area Chair。", "其他服务：", "手写协会会员"])
        section("自定义栏目：", ["不要覆盖这个字段"])
        section("论文著作：", ["完整论文列表请见https://example.test/publications"])
        document.add_paragraph("手写署名与说明")
        target = root / "form.docx"
        document.save(target)
        return target, data

    def test_fill_missing_content_preserve_manual_text_and_package_parts(self):
        with tempfile.TemporaryDirectory() as folder:
            target, data = self.fixture(folder)
            original = target.read_bytes()
            before = Document(target)
            education = etree.tostring(before.tables[0].rows[1]._tr)
            result = update_faculty_form(target, data_dir=data)
            document = Document(target)
            text = "\n".join(row.cells[0].text for row in document.tables[0].rows)
            self.assertIn("教授课程：", text)
            self.assertIn("CSE 188", text)
            self.assertIn("新增基金项目", text)
            self.assertNotIn("旧项目", text)
            self.assertIn("Action Editor", text)
            self.assertIn("Senior Area Chair", text)
            self.assertIn("The Web Conference (WWW)", text)
            self.assertIn("Journal B", text)
            self.assertIn("手写协会会员", text)
            self.assertIn("不要覆盖这个字段", text)
            self.assertIn("任教期间：", text)
            self.assertEqual(etree.tostring(document.tables[0].rows[1]._tr), education)
            self.assertEqual(document.paragraphs[-1].text, "手写署名与说明")
            with ZipFile(result.backup) as archive:
                self.assertEqual(archive.namelist(), [target.name])
                self.assertEqual(archive.read(target.name), original)
            with ZipFile(BytesIO(original)) as old, ZipFile(target) as new:
                for name in old.namelist():
                    if name not in ("word/document.xml", "word/numbering.xml"):
                        self.assertEqual(old.read(name), new.read(name), name)
            self.assertEqual(list(Path(folder).rglob("*.docx")), [target])

    def test_repeated_update_is_a_noop_without_duplicate_rows_or_backups(self):
        with tempfile.TemporaryDirectory() as folder:
            target, data = self.fixture(folder)
            update_faculty_form(target, data_dir=data)
            once = target.read_bytes()
            second = update_faculty_form(target, data_dir=data)
            self.assertFalse(second.changed)
            self.assertEqual(target.read_bytes(), once)
            self.assertEqual(len(list((Path(folder) / "generated/form_backups").glob("*.zip"))), 1)

    def test_locked_target_leaves_original_without_any_word_copies(self):
        with tempfile.TemporaryDirectory() as folder:
            target, data = self.fixture(folder)
            original = target.read_bytes()
            real_replace = os.replace
            def locked(source, destination):
                if Path(destination) == target:
                    raise PermissionError("Word is open")
                return real_replace(source, destination)
            with patch("faculty_form_word.os.replace", side_effect=locked):
                with self.assertRaisesRegex(PermissionError, "请保存并关闭"):
                    update_faculty_form(target, data_dir=data, backups=False)
            self.assertEqual(target.read_bytes(), original)
            self.assertEqual(list(Path(folder).rglob("*.docx")), [target])
            self.assertEqual(list(Path(folder).glob(".faculty_form_*")), [])

    def test_full_publications_are_optional_and_keep_source_labels(self):
        with tempfile.TemporaryDirectory() as folder:
            target, data = self.fixture(folder)
            source = Path(folder) / "publications.md"
            source.write_text('<ol><li><div class="title">A Study</div><div class="author">A. Author.</div><div class="periodical"><em>Example, 2026.</em><span>(CCF-A)</span><span>(JCR Q1)</span></div></li></ol>', encoding="utf-8")
            result = update_faculty_form(target, data_dir=data, publications_source=source, publications_mode="full")
            document = Document(result.path)
            paragraphs = [p for row in document.tables[0].rows for p in row.cells[0].paragraphs]
            paper = next(p for p in paragraphs if "A Study" in p.text)
            labels = [r for r in paper.runs if str(r.font.color.rgb) == "FF0000"]
            self.assertEqual([r.text.replace("\u00a0", " ").replace("\u2011", "-") for r in labels], [" (CCF-A)", " (JCR Q1)"])
            self.assertFalse(any("完整论文列表请见" in p.text for p in paragraphs))
            self.assertIsNotNone(paper._p.pPr.numPr)


if __name__ == "__main__":
    unittest.main()
